from pathlib import Path

import docx
import pymupdf
import pytest

from doc_holler.core.errors import EmptyDocumentError, UnsupportedFormatError
from doc_holler.integrations.document_extractors import (
    extract_document,
    extract_docx,
    extract_pdf,
    extract_txt,
)


def test_extract_txt(tmp_path: Path) -> None:
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Hello Doc Holler world!", encoding="utf-8")
    sections = extract_txt(txt_file)
    assert len(sections) == 1
    assert sections[0].text == "Hello Doc Holler world!"
    assert sections[0].page_number is None


def test_extract_txt_empty(tmp_path: Path) -> None:
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("   \n\n  ", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        extract_txt(empty_file)


def test_extract_pdf(tmp_path: Path) -> None:
    pdf_path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "First page content for Doc Holler")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Second page content for retrieval")
    doc.save(pdf_path)
    doc.close()

    sections = extract_pdf(pdf_path)
    assert len(sections) == 2
    assert "First page" in sections[0].text
    assert sections[0].page_number == 1
    assert "Second page" in sections[1].text
    assert sections[1].page_number == 2


def test_extract_pdf_empty(tmp_path: Path) -> None:
    pdf_path = tmp_path / "empty.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()

    with pytest.raises(EmptyDocumentError):
        extract_pdf(pdf_path)


def test_extract_docx(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_paragraph("Docx introduction paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    doc.save(docx_path)

    sections = extract_docx(docx_path)
    assert len(sections) == 1
    assert "Docx introduction" in sections[0].text
    assert "Header A | Header B" in sections[0].text
    assert sections[0].page_number is None


def test_extract_docx_empty(tmp_path: Path) -> None:
    docx_path = tmp_path / "empty.docx"
    doc = docx.Document()
    doc.save(docx_path)

    with pytest.raises(EmptyDocumentError):
        extract_docx(docx_path)


def test_extract_document_unsupported(tmp_path: Path) -> None:
    unsupported_path = tmp_path / "test.csv"
    unsupported_path.write_text("a,b,c", encoding="utf-8")
    with pytest.raises(UnsupportedFormatError):
        extract_document(unsupported_path, "test.csv")
