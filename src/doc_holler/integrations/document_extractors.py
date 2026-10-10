from dataclasses import dataclass
from pathlib import Path

import docx
import pymupdf

from doc_holler.core.errors import (
    EmptyDocumentError,
    ExtractionError,
    UnsupportedFormatError,
)


@dataclass(frozen=True)
class ExtractedSection:
    text: str
    page_number: int | None = None


def extract_pdf(file_path: Path) -> list[ExtractedSection]:
    try:
        doc = pymupdf.open(file_path)
    except Exception as exc:
        raise ExtractionError(f"Failed to open PDF document: {exc}") from exc

    sections: list[ExtractedSection] = []
    try:
        for page_index in range(len(doc)):
            page = doc[page_index]
            text = page.get_text().strip()
            if text:
                sections.append(ExtractedSection(text=text, page_number=page_index + 1))
    finally:
        doc.close()

    if not sections:
        raise EmptyDocumentError("PDF document contains no extractable text")

    return sections


def extract_docx(file_path: Path) -> list[ExtractedSection]:
    try:
        doc = docx.Document(file_path)
    except Exception as exc:
        raise ExtractionError(f"Failed to open DOCX document: {exc}") from exc

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs).strip()
    if not full_text:
        raise EmptyDocumentError("DOCX document contains no extractable text")

    return [ExtractedSection(text=full_text, page_number=None)]


def extract_txt(file_path: Path) -> list[ExtractedSection]:
    try:
        content = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            content = file_path.read_text(encoding="latin-1")
        except Exception as exc:
            raise ExtractionError(f"Failed to read TXT document: {exc}") from exc
    except Exception as exc:
        raise ExtractionError(f"Failed to read TXT document: {exc}") from exc

    cleaned = content.strip()
    if not cleaned:
        raise EmptyDocumentError("TXT document contains no extractable text")

    return [ExtractedSection(text=cleaned, page_number=None)]


def extract_document(file_path: Path, filename: str) -> list[ExtractedSection]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(file_path)
    if suffix == ".docx":
        return extract_docx(file_path)
    if suffix == ".txt":
        return extract_txt(file_path)
    raise UnsupportedFormatError(f"Unsupported file format '{suffix}'")
